from pathlib import Path
from tempfile import TemporaryDirectory
import logging
from aiogram import Router, F
from bot.config import get_settings
from bot.services.llm import OpenAICompatibleLLMService
from bot.services.tz_generator import TZGenerator
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.fsm.context import FSMContext
from aiogram.filters import StateFilter
from docx import Document
from fpdf import FPDF
from bot.states.dialogue import DialogueStates


router = Router(name=__name__)
logger = logging.getLogger(__name__)


def _save_tz_files(tz_text: str, output_dir: Path) -> tuple[Path, Path]:
    docx_path = output_dir / "technical_specification.docx"
    pdf_path = output_dir / "technical_specification.pdf"

    doc = Document()
    doc.add_heading("Техническое задание", level=0)
    for line in tz_text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        else:
            doc.add_paragraph(line.removeprefix("- ").removeprefix("* "))
    doc.save(docx_path)

    pdf = FPDF(format="A4")
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font_candidates = (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/Library/Fonts/Arial.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    )
    font_path = next((path for path in font_candidates if path.exists()), None)
    if font_path is None:
        raise FileNotFoundError(
            "Не найден Unicode-шрифт для PDF. Установите DejaVu Sans "
            "(пакет fonts-dejavu-core) или добавьте путь к TTF в font_candidates."
        )
    pdf.add_font("Unicode", "", str(font_path))
    pdf.set_font("Unicode", size=11)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(w=pdf.epw, h=8, text="Техническое задание")
    pdf.ln(3)
    for line in tz_text.splitlines():
        line = line.strip()
        if line:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(
                w=pdf.epw,
                h=7,
                text=line.removeprefix("- ").removeprefix("* "),
            )
    pdf.output(str(pdf_path))
    return docx_path, pdf_path

def get_tz_generator() -> TZGenerator:
    settings = get_settings()
    llm = OpenAICompatibleLLMService(
        api_key=settings.llm_api_key.get_secret_value(),
        base_url=settings.llm_base_url,
        model=settings.llm_model,
    )
    return TZGenerator(llm)

@router.callback_query(F.data == "start_interview")
async def start_interview(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await state.set_state(DialogueStates.interviewing)
    await state.update_data(history=[], answers={})

    generator = get_tz_generator()
    first_question = await generator.generate_next_message([
        {"role": "user", "content": "Начни интервью. Задай первый вопрос о проекте."}
    ])

    await state.update_data(history=[{"role": "assistant", "content": first_question}])

    await callback.message.answer(f"👋 Отлично, давайте вместе подготовим ТЗ!\n\n💬 {first_question}")

@router.message(StateFilter(DialogueStates.interviewing))
async def process_answer(message: Message, state: FSMContext):
    try:
        data = await state.get_data()
        history = data.get("history", [])
        answers = data.get("answers", {})

        answer = message.text or ""
        history.append({"role": "user", "content": answer})
        answers[f"answer_{len(answers)+1}"] = answer

        generator = get_tz_generator()
        next_response = await generator.generate_next_message(history)

        if next_response.strip().upper().startswith("DONE"):
            await state.set_state(DialogueStates.generating)
            await state.update_data(history=history, answers=answers)
            await message.answer("✅ Спасибо за ответы! Уже собираю ваше техническое задание 📄✨")

            tz_text = await generator.generate_tz(answers, history)
            with TemporaryDirectory(prefix="telegram-tz-") as temp_dir:
                docx_path, pdf_path = _save_tz_files(tz_text, Path(temp_dir))
                await message.answer(
                    "🎉 Готово! Ваше техническое задание подготовлено.\n"
                    "📎 Прикрепила версии в DOCX и PDF — выберите удобную!"
                )
                await message.answer_document(
                    FSInputFile(docx_path), caption="📝 Техническое задание (DOCX)"
                )
                await message.answer_document(
                    FSInputFile(pdf_path), caption="📄 Техническое задание (PDF)"
                )
            await state.clear()
            return

        history.append({"role": "assistant", "content": next_response})
        await state.update_data(history=history, answers=answers)
        await message.answer(f"💬 {next_response}")
    except Exception:
        logger.exception("Ошибка при обработке ответа или подготовке/отправке ТЗ")
        await state.set_state(DialogueStates.interviewing)
        await message.answer(
            "😔 Не получилось обработать ответ или подготовить файлы. "
            "Ваши ответы сохранены — попробуйте ещё раз чуть позже. "
            "Если ошибка повторится, сообщите администратору бота."
        )
